from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Iterable

from .schemas import DocumentChunk

logger = logging.getLogger("kb_web_agent.documents")

HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")

SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".docx", ".xlsx"}


# ---------------------------------------------------------------------------
# 文件枚举
# ---------------------------------------------------------------------------


def iter_documents(root: Path) -> Iterable[Path]:
    """递归枚举所有支持格式的文档文件（md/txt/pdf/docx/xlsx）。"""
    if root.is_file() and root.suffix.lower() in SUPPORTED_EXTENSIONS:
        yield root
        return
    if not root.exists():
        raise FileNotFoundError(f"Knowledge base path does not exist: {root}")
    if not root.is_dir():
        raise NotADirectoryError(f"Knowledge base path is not a directory: {root}")
    yield from sorted(p for p in root.rglob("*") if p.suffix.lower() in SUPPORTED_EXTENSIONS)


def iter_markdown_files(root: Path) -> Iterable[Path]:
    """向后兼容别名，枚举 .md / .txt 文件。"""
    if root.is_file() and root.suffix.lower() in (".md", ".txt"):
        yield root
        return
    if not root.exists():
        raise FileNotFoundError(f"Knowledge base path does not exist: {root}")
    if not root.is_dir():
        raise NotADirectoryError(f"Knowledge base path is not a directory: {root}")
    yield from sorted(p for p in root.rglob("*") if p.suffix.lower() in (".md", ".txt"))


# ---------------------------------------------------------------------------
# 格式读取
# ---------------------------------------------------------------------------


def read_markdown(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def read_pdf(path: Path) -> str:
    """使用 pdfplumber 提取 PDF 文本，每页以空行分隔。"""
    try:
        import pdfplumber
    except ImportError:
        logger.warning("[Documents] pdfplumber 未安装，跳过 PDF: %s", path)
        return ""
    pages: list[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(text)
    except Exception as exc:
        logger.warning("[Documents] PDF 解析失败 path=%s err=%s", path, exc)
        return ""
    return "\n\n".join(pages)


def read_docx(path: Path) -> str:
    """使用 python-docx 提取 Word 文本，每段以换行分隔。"""
    try:
        from docx import Document
    except ImportError:
        logger.warning("[Documents] python-docx 未安装，跳过 Word: %s", path)
        return ""
    try:
        doc = Document(str(path))
        return "\n".join(para.text for para in doc.paragraphs if para.text.strip())
    except Exception as exc:
        logger.warning("[Documents] DOCX 解析失败 path=%s err=%s", path, exc)
        return ""


def read_excel(path: Path) -> str:
    """使用 openpyxl 提取 Excel 内容，每行转为 tab 分隔字符串。"""
    try:
        import openpyxl
    except ImportError:
        logger.warning("[Documents] openpyxl 未安装，跳过 Excel: %s", path)
        return ""
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            lines.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    lines.append("\t".join(cells))
        return "\n".join(lines)
    except Exception as exc:
        logger.warning("[Documents] Excel 解析失败 path=%s err=%s", path, exc)
        return ""


def read_document(path: Path) -> str:
    """按扩展名分发到对应读取函数，返回纯文本。"""
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return read_markdown(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    if ext == ".xlsx":
        return read_excel(path)
    logger.warning("[Documents] 不支持的格式，跳过: %s", path)
    return ""


# ---------------------------------------------------------------------------
# 分块
# ---------------------------------------------------------------------------


def split_markdown(text: str, path: Path, max_chars: int = 1400) -> list[DocumentChunk]:
    """按 Markdown 标题分块（同时适用于纯文本和其他格式的文本内容）。"""
    lines = text.splitlines()
    chunks: list[DocumentChunk] = []
    heading_stack: list[tuple[int, str]] = []
    buffer: list[str] = []
    start_line = 1
    current_title = path.stem

    def flush(end_line: int) -> None:
        nonlocal buffer, start_line
        content = "\n".join(buffer).strip()
        if content:
            chunks.append(
                DocumentChunk(
                    id=f"{path.name}:{start_line}-{end_line}",
                    path=path,
                    title=current_title,
                    text=content,
                    line_start=start_line,
                    line_end=end_line,
                )
            )
        buffer = []
        start_line = end_line + 1

    for index, line in enumerate(lines, start=1):
        heading = HEADING_RE.match(line)
        if heading:
            if buffer:
                flush(index - 1)
            level = len(heading.group(1))
            title = heading.group(2).strip()
            while heading_stack and heading_stack[-1][0] >= level:
                heading_stack.pop()
            heading_stack.append((level, title))
            current_title = " / ".join(item[1] for item in heading_stack) or path.stem
            start_line = index

        if not buffer:
            start_line = index
        buffer.append(line)
        if sum(len(item) + 1 for item in buffer) >= max_chars and not line.strip():
            flush(index)

    if buffer:
        flush(len(lines) or 1)
    return chunks


# ---------------------------------------------------------------------------
# 批量加载
# ---------------------------------------------------------------------------


def load_document_chunks(root: Path, max_chars: int = 1400) -> list[DocumentChunk]:
    """加载所有支持格式的文档并分块（新接口）。"""
    chunks: list[DocumentChunk] = []
    for path in iter_documents(root):
        text = read_document(path)
        if text:
            chunks.extend(split_markdown(text, path, max_chars=max_chars))
    return chunks


def load_markdown_chunks(root: Path, max_chars: int = 1400) -> list[DocumentChunk]:
    """向后兼容别名，等同于 load_document_chunks。"""
    return load_document_chunks(root, max_chars=max_chars)
